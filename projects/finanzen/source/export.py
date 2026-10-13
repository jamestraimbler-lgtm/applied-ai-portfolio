"""
export.py — exporteren naar Excel
"""
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import database as db

TEAL      = "1F4E5F"
LICHT     = "EEF2F4"
WIT       = "FFFFFF"
GROEN     = "E7F0E8"
ROOD_L    = "FDECEA"


def _header_stijl(ws, rij, kolommen: list[str]):
    for col, titel in enumerate(kolommen, 1):
        c = ws.cell(row=rij, column=col, value=titel)
        c.font      = Font(bold=True, color=WIT, size=11)
        c.fill      = PatternFill("solid", fgColor=TEAL)
        c.alignment = Alignment(wrap_text=True)


def _rand():
    dun = Side(style='thin', color='BBBBBB')
    return Border(left=dun, right=dun, top=dun, bottom=dun)


def _stuk_kleur(status: str) -> str:
    return {'gecontroleerd': GROEN, 'ontvangen': LICHT,
            'ontbreekt': ROOD_L}.get(status, WIT)


def exporteer_dossier(zaak_id: int, dest_map: str) -> str:
    """Exporteer één dossier naar een gedetailleerd Excel-bestand."""
    data  = db.zaak_volledig_ophalen(zaak_id)
    zaak  = data['zaak']
    k1    = data['klant1']
    k2    = data['klant2']

    naam1 = db.klant_volledige_naam(k1)
    naam2 = db.klant_volledige_naam(k2) if k2 else ''
    namen = f"{naam1} & {naam2}" if naam2 else naam1
    ts    = datetime.now().strftime('%Y%m%d')
    safe  = "".join(c for c in namen if c.isalnum() or c in ' _-')[:40]
    bestand = os.path.join(dest_map, f"Dossier_{safe}_{ts}.xlsx")

    wb = openpyxl.Workbook()

    # ── Blad 1: Overzicht ────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Overzicht"
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 42

    ws['A1'] = "FinanZen — Dossierexport"
    ws['A1'].font = Font(bold=True, size=14, color=TEAL)
    ws['A2'] = f"Gegenereerd op {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws['A2'].font = Font(italic=True, color='888888')

    rij = 4
    for lbl, val in [
        ("Aanvrager 1",  naam1),
        ("Aanvrager 2",  naam2 or '—'),
        ("Soort dossier",db.ZAAK_SOORTEN.get(zaak['soort'], zaak['soort'])),
        ("Omschrijving", zaak['omschrijving'] or ''),
        ("Status",       db.ZAAK_STATUSSEN.get(zaak['status'], zaak['status'])),
        ("Aangemaakt",   zaak['created_at']),
        ("Bijgewerkt",   zaak['updated_at']),
    ]:
        ws.cell(row=rij, column=1, value=lbl).font = Font(bold=True)
        ws.cell(row=rij, column=2, value=val)
        rij += 1

    # ── Blad 2: Stukkenlijst ─────────────────────────────────────────────────
    ws2 = wb.create_sheet("Stukkenlijst")
    ws2.column_dimensions['A'].width = 40
    ws2.column_dimensions['B'].width = 20
    ws2.column_dimensions['C'].width = 15
    _header_stijl(ws2, 1, ["Document", "Status", "Verplicht"])
    for i, stuk in enumerate(data['stukken'], 2):
        stlbl = db.STUK_STATUSSEN.get(stuk['status'], (stuk['status'], ''))[0]
        for col, val in [(1, stuk['label']), (2, stlbl),
                         (3, 'Ja' if stuk['verplicht'] else 'Optioneel')]:
            c = ws2.cell(row=i, column=col, value=val)
            c.fill   = PatternFill("solid", fgColor=_stuk_kleur(stuk['status']))
            c.border = _rand()

    # ── Blad 3: Personalia ────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Personalia")
    ws3.column_dimensions['A'].width = 22
    ws3.column_dimensions['B'].width = 30
    ws3.column_dimensions['C'].width = 22
    ws3.column_dimensions['D'].width = 30

    _header_stijl(ws3, 1, ["Aanvrager 1", "", "Aanvrager 2", ""])
    velden = [
        ("Voornaam",     k1.get('voornaam',''), 'voornaam'),
        ("Tussenvoegsel",k1.get('tussenvoegsel',''), 'tussenvoegsel'),
        ("Achternaam",   k1.get('achternaam',''), 'achternaam'),
        ("Geboortedatum",k1.get('geboortedatum',''), 'geboortedatum'),
        ("E-mail",       k1.get('email',''), 'email'),
        ("Telefoon",     k1.get('telefoon',''), 'telefoon'),
        ("Adres",        k1.get('adres',''), 'adres'),
        ("Postcode",     k1.get('postcode',''), 'postcode'),
        ("Woonplaats",   k1.get('woonplaats',''), 'woonplaats'),
    ]
    for i, (lbl, val1, key) in enumerate(velden, 2):
        val2 = k2.get(key, '') if k2 else ''
        ws3.cell(row=i, column=1, value=lbl).font = Font(bold=True)
        ws3.cell(row=i, column=2, value=val1)
        ws3.cell(row=i, column=3, value=lbl).font = Font(bold=True)
        ws3.cell(row=i, column=4, value=val2)

    # ── Blad 4: Taken ────────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Taken")
    ws4.column_dimensions['A'].width = 44
    ws4.column_dimensions['B'].width = 12
    ws4.column_dimensions['C'].width = 16
    _header_stijl(ws4, 1, ["Omschrijving", "Status", "Deadline"])
    for i, taak in enumerate(data['taken'], 2):
        for col, val in [(1, taak['omschrijving']),
                         (2, 'Gedaan' if taak['status'] == 'gedaan' else 'Open'),
                         (3, taak['deadline'] or '')]:
            c = ws4.cell(row=i, column=col, value=val)
            c.border = _rand()

    wb.save(bestand)
    return bestand


def exporteer_alle_dossiers(dest_map: str) -> str:
    """Exporteer een overzicht van alle dossiers naar één Excel-bestand."""
    zaken   = db.alle_zaken_ophalen()
    ts      = datetime.now().strftime('%Y%m%d_%H%M')
    bestand = os.path.join(dest_map, f"FinanZen_alle_dossiers_{ts}.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alle dossiers"
    for breedte, kol in [(5,'A'),(34,'B'),(34,'C'),(22,'D'),(18,'E'),(14,'F'),(12,'G')]:
        ws.column_dimensions[kol].width = breedte

    _header_stijl(ws, 1, ['#', 'Aanvrager 1', 'Aanvrager 2', 'Soort',
                           'Omschrijving', 'Status', 'Ontbrekende stukken'])
    for i, z in enumerate(zaken, 2):
        vals = [z['id'], z['naam1'], z['naam2'] or '',
                db.ZAAK_SOORTEN.get(z['soort'], z['soort']),
                z['omschrijving'] or '',
                db.ZAAK_STATUSSEN.get(z['status'], z['status']),
                z['ontbrekend']]
        for col, val in enumerate(vals, 1):
            c = ws.cell(row=i, column=col, value=val)
            c.border = _rand()
            if z['ontbrekend'] and col == 7:
                c.fill = PatternFill("solid", fgColor="FDECEA")

    ws.freeze_panes = 'A2'
    wb.save(bestand)
    return bestand
